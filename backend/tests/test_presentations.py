from __future__ import annotations

import io
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

import fileora.presentations as presentations
from fileora.chunking import chunk_units
from fileora.domain import Extraction, FileoraError, Unit
from fileora.evaluation import span_matches
from fileora.extraction import extract, supervised_extract
from fileora.presentation_ocr import PresentationOCR
from fileora.retrieval import Filters, SearchRequest


@pytest.fixture(autouse=True)
def portable_test_environment(monkeypatch):
    # Tests mock the portable recognizer independently of host OCR installations.
    original = shutil.which
    monkeypatch.delenv("FILEORA_TESSERACT_CMD", raising=False)
    monkeypatch.setattr(
        shutil, "which", lambda command: None if command == "tesseract" else original(command)
    )


def write_deck(path: Path, keyword="BCNF", image=False):
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[6])
    first.shapes.add_textbox(0, 0, Inches(6), Inches(1)).text = "Database lecture"
    first.notes_slide.notes_text_frame.text = "Instructor reminder: quasarproof"
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(
        0, 0, Inches(6), Inches(1)
    ).text = f"{keyword} normalization removes duplicated relational data."
    table = slide.shapes.add_table(2, 2, 0, Inches(1), Inches(6), Inches(2)).table
    table.cell(0, 0).text = "Determinant"
    table.cell(0, 1).text = "Superkey"
    table.cell(1, 0).text = "Enrollment"
    table.cell(1, 1).text = "Student identifier"
    if image:
        buffer = io.BytesIO()
        Image.new("RGB", (30, 30), "white").save(buffer, format="PNG")
        slide.shapes.add_picture(io.BytesIO(buffer.getvalue()), 0, Inches(3))
    deck.save(path)
    return path


def test_slide_group_table_and_notes_provenance(tmp_path, service):
    path = write_deck(tmp_path / "lecture.pptx")
    result = extract(path, service.settings)
    assert result.modality == "presentation"
    assert [(u.locator["slide"], u.locator["section"]) for u in result.units] == [
        (1, "slide"),
        (1, "notes"),
        (2, "slide"),
    ]
    assert result.units[1].text == "Instructor reminder: quasarproof"
    assert "BCNF" in result.units[2].text
    assert "Determinant | Superkey" in result.units[2].text
    chunks = chunk_units(result.units, 8, 2)
    assert {c.locator["slide"] for c in chunks} == {1, 2}
    assert all("quasarproof" not in c.text for c in chunks if c.locator["section"] == "slide")


def test_powerpoint_search_filter_updates_and_deletes(service, corpus):
    path = write_deck(corpus / "lecture.pptx")
    job = service.indexer.run(service.indexer.create_job())
    assert job["indexed"] == 1 and job["failed"] == 0
    for mode in ("lexical", "semantic", "hybrid"):
        response = service.search.run(
            SearchRequest(
                query="BCNF normalization", mode=mode, filters=Filters(modality="presentation")
            )
        )
        assert response["results"][0]["name"] == path.name
        assert response["results"][0]["evidence"][0]["locator"]["slide"] == 2
    response = service.search.run(SearchRequest(query="quasarproof"))
    assert response["results"][0]["evidence"][0]["locator"]["section"] == "notes"
    for modality in ("presentation", "code", "image", "document"):
        query = "unfindablepptsymbol" if modality == "presentation" else "BCNF"
        assert not service.search.run(
            SearchRequest(query=query, filters=Filters(modality=modality))
        )["results"]
    write_deck(path, keyword="Replacementword")
    service.indexer.run(service.indexer.create_job(verify=True))
    assert not service.search.run(SearchRequest(query="BCNF"))["results"]
    path.unlink()
    service.indexer.run(service.indexer.create_job())
    assert not service.search.run(SearchRequest(query="Replacementword"))["results"]


def test_powerpoint_source_download_and_detail(client, service, corpus):
    path = write_deck(corpus / "lecture.pptx")
    service.indexer.run(service.indexer.create_job())
    response = client.post(
        "/api/v1/search", json={"query": "BCNF", "filters": {"modality": "presentation"}}
    )
    assert response.status_code == 200
    file_id = response.json()["results"][0]["file_id"]
    detail = client.get(f"/api/v1/files/{file_id}").json()
    assert any(c["locator"]["slide"] == 2 for c in detail["chunks"])
    original = client.get(f"/api/v1/files/{file_id}/preview")
    assert original.content == path.read_bytes()
    assert (
        original.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    assert "attachment" in original.headers["content-disposition"]


def test_embedded_image_ocr_preserves_slide_and_asset(tmp_path, service, monkeypatch):
    path = write_deck(tmp_path / "images.pptx", image=True)
    result = extract(path, service.settings)
    assert "slide_2:image_text_requires_ocr" in result.warnings
    service.settings.enable_ocr = True
    monkeypatch.setattr(
        PresentationOCR,
        "_portable",
        lambda self, image, timeout: Unit("imagekeyword", "ocr", {"boxes": []}),
    )
    result = extract(path, service.settings)
    unit = next(u for u in result.units if u.kind == "ocr")
    assert unit.locator["slide"] == 2 and unit.locator["section"] == "image"
    assert (service.settings.data_dir / "assets" / unit.asset).is_file()
    service.settings.max_image_pixels = 100
    result = extract(path, service.settings)
    assert "slide_2:IMAGE_TOO_LARGE" in result.warnings
    assert any("BCNF" in u.text for u in result.units)


def test_scatter_chart_does_not_drop_slide_text(tmp_path, service):
    path = write_deck(tmp_path / "chart.pptx")
    deck = Presentation(path)
    data = XyChartData()
    series = data.add_series("Series label")
    series.add_data_point(1, 2)
    deck.slides[0].shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, 0, 0, Inches(2), Inches(2), data)
    deck.save(path)
    result = extract(path, service.settings)
    assert "Series label" in result.units[0].text
    assert "BCNF" in result.units[-1].text


def test_powerpoint_spawn(tmp_path, service):
    result = supervised_extract(write_deck(tmp_path / "spawn.pptx"), service.settings)
    assert result.modality == "presentation"
    assert result.units[-1].locator["slide"] == 2


def test_empty_and_corrupt_presentations(tmp_path, service):
    path = tmp_path / "empty.pptx"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(path)
    assert extract(path, service.settings).warnings == ["slide_1:no_slide_text"]
    path.write_bytes(b"not a presentation")
    with pytest.raises(FileoraError) as error:
        extract(path, service.settings)
    assert error.value.code == "PRESENTATION_UNREADABLE"


def test_expanded_package_limit(tmp_path, service):
    path = tmp_path / "expanded.pptx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as package:
        package.writestr("ppt/presentation.xml", "x" * 5000)
    service.settings.max_file_bytes = 1000
    with pytest.raises(FileoraError) as error:
        extract(path, service.settings)
    assert error.value.code == "PRESENTATION_TOO_LARGE"


def test_missing_legacy_converter_is_actionable(service, corpus, monkeypatch):
    monkeypatch.setattr(presentations, "libreoffice_command", lambda: None)
    (corpus / "legacy.ppt").write_bytes(b"legacy")
    job = service.indexer.run(service.indexer.create_job())
    assert job["failed"] == 1
    error = service.store.one("SELECT code,message FROM job_errors WHERE job_id=?", (job["id"],))
    assert error["code"] == "LEGACY_PPT_UNAVAILABLE"
    assert ".pptx" in error["message"]
    assert not list(service.settings.data_dir.glob("ppt-convert-*"))


def test_legacy_conversion_contract(tmp_path, service, monkeypatch):
    path = tmp_path / "legacy.ppt"
    path.write_bytes(b"original legacy bytes")
    monkeypatch.setattr(presentations, "libreoffice_command", lambda: "local-soffice")
    seen = []

    class Converter:
        returncode = 0

        def __init__(self, args, **kwargs):
            seen.append(args)
            output = Path(args[args.index("--outdir") + 1])
            security = (
                output.parent / "profile" / "user" / "registrymodifications.xcu"
            ).read_text()
            for name in (
                "DisableMacrosExecution",
                "DisableActiveContent",
                "BlockUntrustedRefererLinks",
            ):
                assert name in security
            assert security.count("<value>true</value>") == 3
            write_deck(output / "legacy.pptx")

        def wait(self, **kwargs):
            return 0

    monkeypatch.setattr(presentations.subprocess, "Popen", Converter)
    result = extract(path, service.settings)
    assert result.units[-1].locator["slide"] == 2
    assert "legacy_ppt_converted_locally" in result.warnings
    assert path.read_bytes() == b"original legacy bytes"
    assert seen[0][-1] == str(path.resolve())
    assert seen[0][1].startswith("-env:UserInstallation=file:")
    assert not list(service.settings.data_dir.glob("ppt-convert-*"))


def test_converter_timeout(tmp_path, service, monkeypatch):
    monkeypatch.setattr(presentations, "libreoffice_command", lambda: "local-soffice")
    monkeypatch.setattr(presentations.subprocess, "run", lambda *args, **kwargs: None)
    stopped = []

    class StalledConverter:
        pid = 12345

        def __init__(self, *args, **kwargs):
            pass

        def wait(self, **kwargs):
            if not stopped:
                raise subprocess.TimeoutExpired("local-soffice", 1)

        def kill(self):
            stopped.append(True)

    monkeypatch.setattr(presentations.subprocess, "Popen", StalledConverter)
    path = tmp_path / "legacy.ppt"
    path.write_bytes(b"legacy")
    with pytest.raises(FileoraError) as error:
        extract(path, service.settings)
    assert error.value.code == "PPT_CONVERSION_TIMEOUT"
    assert stopped and not list(service.settings.data_dir.glob("ppt-convert-*"))


def test_slide_span_evaluation():
    actual = {"slide": 2, "section": "notes"}
    assert span_matches(actual, {"slide": 2, "section": "notes"})
    assert not span_matches(actual, {"slide": 3})
    assert not span_matches(actual, {"slide": 2, "section": "slide"})


def test_ocr_budget_preserves_all_native_slides(tmp_path, service):
    path = write_deck(tmp_path / "bounded.pptx", image=True)
    service.settings.enable_ocr = True
    service.settings.ppt_ocr_seconds = 0
    result = extract(path, service.settings)
    assert any("BCNF" in u.text and u.locator["slide"] == 2 for u in result.units)
    assert any(u.locator["section"] == "notes" for u in result.units)
    assert any(w.startswith("presentation_ocr_partial:") for w in result.warnings)


def test_duplicate_images_ocr_once_with_distinct_slide_locators(tmp_path, service, monkeypatch):
    path = write_deck(tmp_path / "duplicates.pptx", image=True)
    deck = Presentation(path)
    picture = deck.slides[1].shapes[-1]
    deck.slides[0].shapes.add_picture(io.BytesIO(picture.image.blob), 0, Inches(3))
    deck.save(path)
    calls = []

    def recognize(self, image, timeout):
        calls.append(image.size)
        return Unit("imageterm", "ocr", {"boxes": []})

    service.settings.enable_ocr = True
    monkeypatch.setattr(PresentationOCR, "_portable", recognize)
    result = extract(path, service.settings)
    assert len(calls) == 1
    assert {u.locator["slide"] for u in result.units if u.kind == "ocr"} == {1, 2}
    result = extract(path, service.settings)
    assert len(calls) == 1  # Persisted cache also avoids starting a second worker.
    assert {u.locator["slide"] for u in result.units if u.kind == "ocr"} == {1, 2}


def test_unsupported_image_part_is_warning_not_deck_failure(tmp_path, service):
    path = write_deck(tmp_path / "unsupported.pptx", image=True)
    # Reproduce a picture whose relationship resolves to a generic OPC Part.
    with zipfile.ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b"image/png", b"application/octet-stream"
    )
    with zipfile.ZipFile(path, "w") as package:
        for name, blob in parts.items():
            package.writestr(name, blob)
    service.settings.enable_ocr = True
    result = extract(path, service.settings)
    assert "slide_2:IMAGE_UNREADABLE" in result.warnings
    assert any("BCNF" in u.text for u in result.units)


def test_ocr_timeout_is_partial_success(tmp_path, service, monkeypatch):
    path = write_deck(tmp_path / "timeout.pptx", image=True)
    service.settings.enable_ocr = True

    def timeout(*args, **kwargs):
        raise FileoraError("OCR_TIMEOUT", "Image OCR timed out")

    monkeypatch.setattr(PresentationOCR, "recognize", timeout)
    result = extract(path, service.settings)
    assert "slide_2:OCR_TIMEOUT" in result.warnings
    assert any("BCNF" in u.text for u in result.units)


def test_ocr_cache_cleanup_when_source_is_forgotten(service, corpus, monkeypatch):
    write_deck(corpus / "images.pptx", image=True)
    service.settings.enable_ocr = True
    monkeypatch.setattr(
        PresentationOCR,
        "_portable",
        lambda *args, **kwargs: Unit("imagekeyword", "ocr", {"boxes": []}),
    )
    service.indexer.run(service.indexer.create_job())
    cache = service.settings.data_dir / "ocr-cache"
    assert list(cache.glob("ppt-*.json"))
    root = service.store.one("SELECT id FROM roots")
    service.indexer.forget_root(root["id"])
    assert not list(cache.glob("ppt-*.json"))


def test_disabled_media_skips_without_parsing_or_errors(service, corpus, monkeypatch):
    (corpus / "lecture.mp4").write_bytes(b"not actually media")
    (corpus / "lecture.wav").write_bytes(b"not actually media")
    job = service.indexer.run(service.indexer.create_job())
    assert job["skipped"] == 4 and job["failed"] == 0
    assert not service.store.rows("SELECT * FROM job_errors WHERE job_id=?", (job["id"],))
    # Enabling transcription still takes the media extraction path.
    service.settings.enable_media = True
    job = service.indexer.run(service.indexer.create_job())
    assert job["failed"] == 2


def test_new_presentation_pipeline_does_not_reindex_other_types(service, corpus):
    plain = service.indexer.pipeline_hash(corpus / "graphs.txt")
    deck = service.indexer.pipeline_hash(corpus / "lecture.pptx")
    service.settings.ppt_ocr_seconds = 10
    assert service.indexer.pipeline_hash(corpus / "graphs.txt") == plain
    assert service.indexer.pipeline_hash(corpus / "lecture.pptx") != deck


def test_duplicate_passages_embed_once_but_keep_each_source_span(service, corpus):
    (corpus / "duplicate.txt").write_text("source fixture")
    service.indexer.extractor = lambda path, settings: Extraction(
        [
            Unit("Repeated lecture logo", locator={"slide": 1}),
            Unit("Repeated lecture logo", locator={"slide": 2}),
        ]
    )
    calls = service.models.calls
    service.indexer.run(service.indexer.create_job())
    assert service.models.calls == calls + 1
    rows = service.store.rows(
        "SELECT c.locator FROM chunks c JOIN file_revisions v ON v.id=c.revision_id JOIN files f ON f.id=v.file_id WHERE f.name='duplicate.txt'"
    )
    assert len(rows) == 2
